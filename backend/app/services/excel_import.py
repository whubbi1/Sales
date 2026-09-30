# backend/app/services/excel_import.py
# Shared Excel helpers: header-row → dict-rows parsing (used by mass_upload.py's generic
# CRM importer and project_management.py's register importers) and a small WHUBBI-branded
# template generator (bold white header row on the app's own teal, used for every
# downloadable "standard template").
import io
from datetime import date, datetime


def _cell_to_json(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def parse_xlsx(content: bytes):
    """First worksheet only. Row 1 is the header; blank rows are skipped. Returns
    (headers, rows) where each row is {header: value}."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [str(h).strip() if h is not None else f"Column {i + 1}" for i, h in enumerate(header_row)]
    rows = []
    for r in rows_iter:
        if r is None or all(c is None for c in r):
            continue
        rows.append({headers[i]: _cell_to_json(r[i]) if i < len(r) else None for i in range(len(headers))})
    return headers, rows


def write_template_xlsx(columns: list[str], sheet_title: str = "Template") -> bytes:
    """A blank, brand-styled starting point: one header row, WHUBBI teal fill/white bold
    text, autosized columns. No data rows — just the columns a register import expects."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel sheet-name length limit

    header_fill = PatternFill(start_color="FF156082", end_color="FF156082", fill_type="solid")
    header_font = Font(color="FFFFFFFF", bold=True, name="Calibri", size=11)
    for col_idx, label in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, min(32, len(label) + 4))
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
