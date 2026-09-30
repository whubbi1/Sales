# backend/app/routers/project_management.py
# Operations > Project Management. Lists exactly the projects Projects Follow-Up lists
# (reuses its list endpoint), then adds per-project governance: basic information and
# templates, members with per-section access, planning (phases nested up to
# PHASE_MAX_DEPTH), tasks, meetings with AI-generated minutes and a validation round, the
# action / risk / decision registers those meetings feed, and versioned deliverables with
# approvals.
#
# Access model: the operations.project_management permission is required for everything.
# 'edit' on it makes the caller a project manager on every project (full access, and the
# only role that can manage members, since members carry the access grants themselves).
# 'view' users only reach a project they're a member of, with the per-section levels from
# their PMMember.permissions.
import asyncio
import io
import json
import os
from datetime import datetime
from typing import List
from uuid import UUID

import boto3
import httpx
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Form, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select, delete, func, text
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.authz import get_current_user_email, access_mode, is_excluded, _LEVELS
from app.services.portal_auth import try_portal_user_email
from app.models.project import Project
from app.models.opportunity import Opportunity
from app.models.company import Company
from app.models.project_management import (
    PM_SECTIONS, DEFAULT_ACTION_STATUSES, TEMPLATE_TYPES, PHASE_MAX_DEPTH,
    PMSettings, PMTemplate, PMMember, PMPhase, PMTask, PMMeeting, PMMeetingValidation,
    PMAction, PMRisk, PMDecision, PMMeetingLink,
    PMDeliverable, PMDeliverableVersion, PMDeliverableApproval,
)
from app.schemas.schemas import ProjectResponse
from app.schemas.project_management import (
    SettingsUpdate, SettingsResponse, TemplateResponse, MemberIn, MemberResponse,
    PhaseIn, PhaseResponse, TaskIn, TaskResponse,
    ActionIn, ActionResponse, RiskIn, RiskResponse, DecisionIn, DecisionResponse,
    MeetingIn, MeetingSummary, MeetingResponse, MeetingReview, Proposal,
    ValidationRequest, ValidationDecision,
    DeliverableIn, DeliverableResponse, VersionIn, AccessResponse,
)
from app.routers.hr import upload_to_s3, s3_ref_to_presigned, AWS_REGION
from app.routers import projects as projects_router

router = APIRouter()

ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY", "")
MINUTES_MODEL = "claude-opus-5"
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_TRANSCRIPT_CHARS = 1_500_000


# ─── Access ──────────────────────────────────────────────────────────────────
class PMUser:
    def __init__(self, email: str, module_mode: str, is_portal: bool = False):
        self.email = email
        self.is_manager = module_mode == "edit"
        self.is_portal = is_portal


async def _contact_has_operation_subscription(db: AsyncSession, email: str) -> bool:
    """The 'Operation' subscription is mandatory for a portal contact to use Project
    Management — a linked Contact without it is treated as having no access at all,
    same as having no PMMember row."""
    r = await db.execute(text("""
        SELECT 1 FROM portal_users pu JOIN contacts c ON c.id = pu.contact_id
        WHERE pu.email = :email AND pu.portal_type = 'partner' AND pu.status = 'active'
        AND c.subscriptions @> '["Operation"]'::jsonb
    """), {"email": email})
    return r.fetchone() is not None


async def pm_user(request: Request, db: AsyncSession = Depends(get_db)) -> PMUser:
    # Portal contacts (added as a PMMember) reach this the same way employees do, just
    # authenticated via the portal's Cognito bearer token instead of X-User-Email — they
    # never carry the internal operations.project_management permission, so that check is
    # skipped entirely for them; every actual read/write still goes through _require()
    # below, which resolves their access purely from their PMMember.permissions.
    portal_email = await try_portal_user_email(request, db)
    if portal_email:
        if not await _contact_has_operation_subscription(db, portal_email):
            raise HTTPException(403, "Your account needs the Operation subscription to use Project Management")
        return PMUser(portal_email, "view", is_portal=True)

    email = await get_current_user_email(request.headers.get("x-user-email"))
    if await is_excluded(email, db):
        raise HTTPException(403, "Access excluded")
    mode = await access_mode(email, "operations", "project_management", db)
    if _LEVELS.get(mode, 0) < _LEVELS["view"]:
        raise HTTPException(403, "No view access to operations.project_management")
    return PMUser(email, mode)


async def _get_project(db: AsyncSession, project_id: UUID) -> Project:
    proj = (await db.execute(select(Project).where(Project.id == project_id))).scalar_one_or_none()
    if not proj:
        raise HTTPException(404, "Project not found")
    return proj


async def _sections_for(db: AsyncSession, project_id: UUID, user: PMUser) -> dict:
    if user.is_manager:
        return {s: "edit" for s in PM_SECTIONS}
    member = (await db.execute(select(PMMember).where(PMMember.project_id == project_id, PMMember.email == user.email))).scalar_one_or_none()
    perms = (member.permissions if member else None) or {}
    sections = {s: perms.get(s, "none") for s in PM_SECTIONS}
    # Members carry access grants, so only managers ever write them (see module header).
    if sections["members"] == "edit":
        sections["members"] = "view"
    return sections


async def _require(db: AsyncSession, project_id: UUID, user: PMUser, section: str, need: str = "view") -> Project:
    proj = await _get_project(db, project_id)
    level = (await _sections_for(db, project_id, user))[section]
    if _LEVELS[level] < _LEVELS[need]:
        raise HTTPException(403, f"No {need} access to {section} on this project")
    return proj


def _require_manager(user: PMUser):
    if not user.is_manager:
        raise HTTPException(403, "Only project managers (edit access to Project Management) can do this")


# ─── Settings & numbering ────────────────────────────────────────────────────
async def _settings(db: AsyncSession, project_id: UUID, lock: bool = False) -> PMSettings:
    # Lazily created on first use; Core applies the model's column defaults to the insert.
    await db.execute(pg_insert(PMSettings).values(project_id=project_id).on_conflict_do_nothing(index_elements=["project_id"]))
    q = select(PMSettings).where(PMSettings.project_id == project_id)
    if lock:
        q = q.with_for_update()
    return (await db.execute(q.execution_options(populate_existing=True))).scalar_one()


async def _next_number(db: AsyncSession, project_id: UUID, prefix: str) -> str:
    s = await _settings(db, project_id, lock=True)
    n = int((s.counters or {}).get(prefix, 0)) + 1
    s.counters = {**(s.counters or {}), prefix: n}
    return f"{prefix}-{n:03d}"


def _action_statuses(s: PMSettings) -> list[str]:
    return DEFAULT_ACTION_STATUSES + [x for x in (s.extra_action_statuses or []) if x not in DEFAULT_ACTION_STATUSES]


async def _project_company_logo(db: AsyncSession, project_id: UUID) -> str | None:
    """A project with no logo of its own falls back to its linked customer's logo."""
    proj = (await db.execute(select(Project).where(Project.id == project_id))).scalar_one_or_none()
    if not proj or not proj.opportunity_id:
        return None
    opp = (await db.execute(select(Opportunity).where(Opportunity.id == proj.opportunity_id))).scalar_one_or_none()
    if not opp or not opp.company_id:
        return None
    company = (await db.execute(select(Company).where(Company.id == opp.company_id))).scalar_one_or_none()
    if not company or not company.logo_url:
        return None
    return await s3_ref_to_presigned(company.logo_url) if company.logo_url.startswith("s3://") else company.logo_url


async def _settings_response(s: PMSettings, fallback_logo_url: str | None = None) -> SettingsResponse:
    return SettingsResponse(
        project_id=s.project_id, sharepoint_url=s.sharepoint_url,
        customer_logo_url=(await s3_ref_to_presigned(s.customer_logo_ref) if s.customer_logo_ref else fallback_logo_url),
        has_custom_logo=bool(s.customer_logo_ref),
        action_statuses=_action_statuses(s), extra_action_statuses=s.extra_action_statuses or [],
        impact_levels=s.impact_levels, probability_levels=s.probability_levels, meeting_types=s.meeting_types,
    )


# ─── S3 read-back (hr.py only offers put + presign) ──────────────────────────
async def _s3_get(ref: str) -> bytes:
    bucket, _, key = ref[5:].partition("/")
    loop = asyncio.get_running_loop()
    obj = await loop.run_in_executor(None, lambda: boto3.client("s3", region_name=AWS_REGION).get_object(Bucket=bucket, Key=key))
    return await loop.run_in_executor(None, obj["Body"].read)


async def _read_upload(file: UploadFile) -> bytes:
    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "File larger than 20 MB")
    return content


def _safe_name(name: str | None) -> str:
    return "".join(c if c.isalnum() or c in "._-" else "_" for c in (name or "file"))[:120]


# ─── Projects list (same set as Projects Follow-Up) ──────────────────────────
@router.get("/projects", response_model=List[ProjectResponse])
async def list_pm_projects(db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    if user.is_portal:
        # Never the full company list for a portal contact — only the project(s) they've
        # been added to as a PMMember, same is_internal exclusion as the employee listing.
        pids = (await db.execute(select(PMMember.project_id).where(PMMember.email == user.email).distinct())).scalars().all()
        if not pids:
            return []
        projects = (await db.execute(
            select(Project).where(Project.id.in_(pids), Project.is_internal.is_(False)).order_by(Project.project_name)
        )).scalars().all()
        await projects_router._attach_related(db, projects)
        return projects
    projects = await projects_router.list_projects(skip=0, limit=500, search=None, is_internal=False, db=db, _="view")
    return projects


async def _require_some_access(db: AsyncSession, project_id: UUID, user: PMUser) -> None:
    # Portal contacts must never be able to probe an arbitrary project id they aren't a
    # member of by guessing a UUID; internal users keep their existing broader visibility.
    if user.is_portal and not any(v != "none" for v in (await _sections_for(db, project_id, user)).values()):
        raise HTTPException(404, "Project not found")


@router.get("/projects/{project_id}", response_model=ProjectResponse)
async def get_pm_project(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require_some_access(db, project_id, user)
    return await projects_router.get_project(project_id, db=db, _="view")


@router.get("/projects/{project_id}/access", response_model=AccessResponse)
async def get_access(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _get_project(db, project_id)
    await _require_some_access(db, project_id, user)
    return AccessResponse(is_manager=user.is_manager, sections=await _sections_for(db, project_id, user))


# ─── Basic information ───────────────────────────────────────────────────────
@router.get("/projects/{project_id}/settings", response_model=SettingsResponse)
async def get_settings(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    # Every section needs the configured statuses/levels/meeting types, so any access at
    # all to the project is enough to read them.
    await _get_project(db, project_id)
    if not any(v != "none" for v in (await _sections_for(db, project_id, user)).values()):
        raise HTTPException(403, "No access to this project")
    s = await _settings(db, project_id)
    await db.commit()
    return await _settings_response(s, await _project_company_logo(db, project_id))


@router.put("/projects/{project_id}/settings", response_model=SettingsResponse)
async def update_settings(project_id: UUID, data: SettingsUpdate, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "basic_info", "edit")
    s = await _settings(db, project_id, lock=True)
    patch = data.model_dump(exclude_unset=True, mode="json")
    for levels_key in ("impact_levels", "probability_levels"):
        if levels_key in patch:
            names = [l["level"].strip() for l in patch[levels_key]]
            if not names or any(not n for n in names) or len(set(names)) != len(names):
                raise HTTPException(400, f"{levels_key}: levels must be non-empty and unique")
    for list_key in ("extra_action_statuses", "meeting_types"):
        if list_key in patch:
            patch[list_key] = list(dict.fromkeys(x.strip() for x in patch[list_key] if x.strip()))
    for k, v in patch.items():
        setattr(s, k, v)
    await db.commit()
    return await _settings_response(s, await _project_company_logo(db, project_id))


@router.post("/projects/{project_id}/logo", response_model=SettingsResponse)
async def upload_logo(project_id: UUID, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "basic_info", "edit")
    if file.content_type not in ("image/png", "image/jpeg"):
        raise HTTPException(400, "Logo must be a PNG or JPEG image")
    content = await _read_upload(file)
    s = await _settings(db, project_id, lock=True)
    s.customer_logo_ref = await upload_to_s3(f"projects/{project_id}/pm/logo", content, file.content_type)
    await db.commit()
    return await _settings_response(s, await _project_company_logo(db, project_id))


@router.delete("/projects/{project_id}/logo", response_model=SettingsResponse)
async def delete_logo(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "basic_info", "edit")
    s = await _settings(db, project_id, lock=True)
    if s.customer_logo_ref and s.customer_logo_ref.startswith("s3://"):
        bucket, _, key = s.customer_logo_ref[5:].partition("/")
        try:
            loop = asyncio.get_running_loop()
            await loop.run_in_executor(None, lambda: boto3.client("s3", region_name=AWS_REGION).delete_object(Bucket=bucket, Key=key))
        except Exception as e:
            print(f"S3 delete error: {e}")
    s.customer_logo_ref = None
    await db.commit()
    return await _settings_response(s, await _project_company_logo(db, project_id))


@router.get("/projects/{project_id}/templates", response_model=List[TemplateResponse])
async def list_templates(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "basic_info")
    r = await db.execute(select(PMTemplate).where(PMTemplate.project_id == project_id).order_by(PMTemplate.template_type, PMTemplate.created_at))
    return r.scalars().all()


@router.post("/projects/{project_id}/templates", response_model=TemplateResponse, status_code=status.HTTP_201_CREATED)
async def upload_template(project_id: UUID, template_type: str = Form(...), name: str = Form(...), file: UploadFile = File(...),
                          db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "basic_info", "edit")
    if template_type not in TEMPLATE_TYPES:
        raise HTTPException(400, f"template_type must be one of {TEMPLATE_TYPES}")
    content = await _read_upload(file)
    tpl = PMTemplate(project_id=project_id, template_type=template_type, name=name.strip() or file.filename,
                     filename=file.filename, content_type=file.content_type, uploaded_by=user.email, file_ref="")
    db.add(tpl)
    await db.flush()
    tpl.file_ref = await upload_to_s3(f"projects/{project_id}/pm/templates/{tpl.id}/{_safe_name(file.filename)}", content,
                                      file.content_type or "application/octet-stream")
    await db.commit()
    return tpl


@router.get("/projects/{project_id}/templates/{template_id}/download")
async def download_template(project_id: UUID, template_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "basic_info")
    tpl = (await db.execute(select(PMTemplate).where(PMTemplate.id == template_id, PMTemplate.project_id == project_id))).scalar_one_or_none()
    if not tpl:
        raise HTTPException(404, "Template not found")
    return {"url": await s3_ref_to_presigned(tpl.file_ref)}


@router.delete("/projects/{project_id}/templates/{template_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_template(project_id: UUID, template_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "basic_info", "edit")
    await db.execute(delete(PMTemplate).where(PMTemplate.id == template_id, PMTemplate.project_id == project_id))
    await db.commit()


# ─── Members ─────────────────────────────────────────────────────────────────
def _clean_member(data: MemberIn) -> dict:
    d = data.model_dump()
    d["email"] = d["email"].strip().lower()
    d["name"] = d["name"].strip()
    if not d["email"] or not d["name"]:
        raise HTTPException(400, "Name and email are required")
    d["permissions"] = {s: d["permissions"].get(s, "none") for s in PM_SECTIONS}
    if d["permissions"]["members"] == "edit":
        d["permissions"]["members"] = "view"
    return d


@router.get("/projects/{project_id}/members", response_model=List[MemberResponse])
async def list_members(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "members")
    r = await db.execute(select(PMMember).where(PMMember.project_id == project_id).order_by(PMMember.name))
    return r.scalars().all()


@router.post("/projects/{project_id}/members", response_model=MemberResponse, status_code=status.HTTP_201_CREATED)
async def add_member(project_id: UUID, data: MemberIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    _require_manager(user)
    await _get_project(db, project_id)
    d = _clean_member(data)
    exists = (await db.execute(select(PMMember.id).where(PMMember.project_id == project_id, PMMember.email == d["email"]))).first()
    if exists:
        raise HTTPException(409, "This email is already a member of the project")
    m = PMMember(project_id=project_id, **d)
    db.add(m)
    await db.commit()
    return m


@router.put("/projects/{project_id}/members/{member_id}", response_model=MemberResponse)
async def update_member(project_id: UUID, member_id: UUID, data: MemberIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    _require_manager(user)
    m = (await db.execute(select(PMMember).where(PMMember.id == member_id, PMMember.project_id == project_id))).scalar_one_or_none()
    if not m:
        raise HTTPException(404, "Member not found")
    d = _clean_member(data)
    if d["email"] != m.email:
        clash = (await db.execute(select(PMMember.id).where(PMMember.project_id == project_id, PMMember.email == d["email"]))).first()
        if clash:
            raise HTTPException(409, "This email is already a member of the project")
    for k, v in d.items():
        setattr(m, k, v)
    await db.commit()
    return m


@router.delete("/projects/{project_id}/members/{member_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_member(project_id: UUID, member_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    _require_manager(user)
    await db.execute(delete(PMMember).where(PMMember.id == member_id, PMMember.project_id == project_id))
    await db.commit()


# ─── Planning ────────────────────────────────────────────────────────────────
async def _check_phase_tree(db: AsyncSession, project_id: UUID, phase_id: UUID | None, parent_id: UUID | None):
    """Rejects a parent from another project, a cycle, or a tree deeper than PHASE_MAX_DEPTH
    once `phase_id` (None = a new leaf) sits under `parent_id`."""
    if parent_id is None:
        parent_depth = 0
        phases = {}
    else:
        rows = (await db.execute(select(PMPhase.id, PMPhase.parent_id).where(PMPhase.project_id == project_id))).all()
        phases = {r.id: r.parent_id for r in rows}
        if parent_id not in phases:
            raise HTTPException(400, "Parent phase not found in this project")
        parent_depth, cur = 0, parent_id
        while cur is not None:
            if cur == phase_id:
                raise HTTPException(400, "A phase can't be moved under itself or one of its sub-phases")
            parent_depth += 1
            cur = phases.get(cur)
    subtree_height = 1
    if phase_id is not None:
        if not phases:
            rows = (await db.execute(select(PMPhase.id, PMPhase.parent_id).where(PMPhase.project_id == project_id))).all()
            phases = {r.id: r.parent_id for r in rows}
        children: dict = {}
        for pid, par in phases.items():
            children.setdefault(par, []).append(pid)

        def height(n):
            return 1 + max((height(c) for c in children.get(n, [])), default=0)
        subtree_height = height(phase_id)
    if parent_depth + subtree_height > PHASE_MAX_DEPTH:
        raise HTTPException(400, f"Planning is limited to {PHASE_MAX_DEPTH} levels")


@router.get("/projects/{project_id}/phases", response_model=List[PhaseResponse])
async def list_phases(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "planning")
    r = await db.execute(select(PMPhase).where(PMPhase.project_id == project_id).order_by(PMPhase.position, PMPhase.start_date, PMPhase.created_at))
    return r.scalars().all()


@router.post("/projects/{project_id}/phases", response_model=PhaseResponse, status_code=status.HTTP_201_CREATED)
async def add_phase(project_id: UUID, data: PhaseIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "planning", "edit")
    await _check_phase_tree(db, project_id, None, data.parent_id)
    if data.position is None or data.position == 0:
        siblings = await db.execute(select(func.coalesce(func.max(PMPhase.position), 0)).where(PMPhase.project_id == project_id, PMPhase.parent_id == data.parent_id if data.parent_id else PMPhase.parent_id.is_(None)))
        data.position = siblings.scalar() + 1
    ph = PMPhase(project_id=project_id, **data.model_dump())
    db.add(ph)
    await db.commit()
    return ph


@router.put("/projects/{project_id}/phases/{phase_id}", response_model=PhaseResponse)
async def update_phase(project_id: UUID, phase_id: UUID, data: PhaseIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "planning", "edit")
    ph = (await db.execute(select(PMPhase).where(PMPhase.id == phase_id, PMPhase.project_id == project_id))).scalar_one_or_none()
    if not ph:
        raise HTTPException(404, "Phase not found")
    if data.parent_id != ph.parent_id:
        await _check_phase_tree(db, project_id, phase_id, data.parent_id)
    for k, v in data.model_dump().items():
        setattr(ph, k, v)
    await db.commit()
    return ph


@router.delete("/projects/{project_id}/phases/{phase_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_phase(project_id: UUID, phase_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "planning", "edit")
    # Sub-phases go with it (FK ON DELETE CASCADE); tasks just lose their phase (SET NULL).
    await db.execute(delete(PMPhase).where(PMPhase.id == phase_id, PMPhase.project_id == project_id))
    await db.commit()


# ─── Tasks ───────────────────────────────────────────────────────────────────
async def _check_task_phase(db, project_id, phase_id):
    if phase_id and not (await db.execute(select(PMPhase.id).where(PMPhase.id == phase_id, PMPhase.project_id == project_id))).first():
        raise HTTPException(400, "Phase not found in this project")


@router.get("/projects/{project_id}/tasks", response_model=List[TaskResponse])
async def list_tasks(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "tasks")
    r = await db.execute(select(PMTask).where(PMTask.project_id == project_id).order_by(PMTask.number))
    return r.scalars().all()


@router.post("/projects/{project_id}/tasks", response_model=TaskResponse, status_code=status.HTTP_201_CREATED)
async def add_task(project_id: UUID, data: TaskIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "tasks", "edit")
    await _check_task_phase(db, project_id, data.phase_id)
    t = PMTask(project_id=project_id, number=await _next_number(db, project_id, "T"), **data.model_dump())
    db.add(t)
    await db.commit()
    return t


@router.put("/projects/{project_id}/tasks/{task_id}", response_model=TaskResponse)
async def update_task(project_id: UUID, task_id: UUID, data: TaskIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "tasks", "edit")
    t = (await db.execute(select(PMTask).where(PMTask.id == task_id, PMTask.project_id == project_id))).scalar_one_or_none()
    if not t:
        raise HTTPException(404, "Task not found")
    await _check_task_phase(db, project_id, data.phase_id)
    for k, v in data.model_dump().items():
        setattr(t, k, v)
    await db.commit()
    return t


@router.delete("/projects/{project_id}/tasks/{task_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_task(project_id: UUID, task_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "tasks", "edit")
    await db.execute(delete(PMTask).where(PMTask.id == task_id, PMTask.project_id == project_id))
    await db.commit()


# ─── Registers: actions / risks / decisions ──────────────────────────────────
# One definition per register; the CRUD routes below are generated from it so the three
# registers can't drift apart in numbering, meeting links or access checks.
REGISTERS = {
    "actions":   dict(model=PMAction,   schema_in=ActionIn,   schema_out=ActionResponse,   prefix="A", item_type="action",   order=PMAction.number),
    "risks":     dict(model=PMRisk,     schema_in=RiskIn,     schema_out=RiskResponse,     prefix="R", item_type="risk",     order=PMRisk.number),
    "decisions": dict(model=PMDecision, schema_in=DecisionIn, schema_out=DecisionResponse, prefix="D", item_type="decision", order=PMDecision.number),
}


async def _links_for(db: AsyncSession, item_type: str, item_ids: list) -> dict:
    if not item_ids:
        return {}
    r = await db.execute(select(PMMeetingLink.item_id, PMMeetingLink.meeting_id).where(PMMeetingLink.item_type == item_type, PMMeetingLink.item_id.in_(item_ids)))
    out: dict = {}
    for item_id, meeting_id in r.all():
        out.setdefault(item_id, []).append(meeting_id)
    return out


async def _set_links(db: AsyncSession, project_id: UUID, item_type: str, item_id: UUID, meeting_ids: list):
    meeting_ids = list(dict.fromkeys(meeting_ids))
    if meeting_ids:
        valid = set((await db.execute(select(PMMeeting.id).where(PMMeeting.project_id == project_id, PMMeeting.id.in_(meeting_ids)))).scalars())
        if len(valid) != len(meeting_ids):
            raise HTTPException(400, "Linked meeting not found in this project")
    await db.execute(delete(PMMeetingLink).where(PMMeetingLink.item_type == item_type, PMMeetingLink.item_id == item_id))
    for mid in meeting_ids:
        db.add(PMMeetingLink(meeting_id=mid, item_type=item_type, item_id=item_id))


async def _add_link(db: AsyncSession, item_type: str, item_id: UUID, meeting_id: UUID):
    await db.execute(pg_insert(PMMeetingLink).values(meeting_id=meeting_id, item_type=item_type, item_id=item_id)
                     .on_conflict_do_nothing(constraint="uq_pm_meeting_link"))


async def _validate_register_fields(db: AsyncSession, project_id: UUID, register: str, values: dict):
    s = await _settings(db, project_id)
    if register == "actions":
        if values.get("status") not in _action_statuses(s):
            raise HTTPException(400, f"Status must be one of {_action_statuses(s)}")
        if values.get("status") == "Solved" and not values.get("closing_date"):
            values["closing_date"] = datetime.utcnow()
        if not values.get("opening_date"):
            values["opening_date"] = datetime.utcnow()
    elif register == "risks":
        for field, levels in (("impact", s.impact_levels), ("probability", s.probability_levels)):
            allowed = [l["level"] for l in levels]
            if values.get(field) not in allowed:
                raise HTTPException(400, f"{field.capitalize()} must be one of {allowed}")


async def _serialize(db: AsyncSession, cfg: dict, items: list) -> list:
    links = await _links_for(db, cfg["item_type"], [i.id for i in items])
    return [cfg["schema_out"].model_validate(i).model_copy(update={"meeting_ids": links.get(i.id, [])}) for i in items]


def _register_routes(register: str, cfg: dict):
    Model, SchemaIn, SchemaOut = cfg["model"], cfg["schema_in"], cfg["schema_out"]
    base = f"/projects/{{project_id}}/{register}"

    async def list_items(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
        await _require(db, project_id, user, register)
        items = (await db.execute(select(Model).where(Model.project_id == project_id).order_by(cfg["order"]))).scalars().all()
        return await _serialize(db, cfg, items)

    async def create_item(project_id: UUID, data: SchemaIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
        await _require(db, project_id, user, register, "edit")
        values = data.model_dump(exclude={"meeting_ids"})
        await _validate_register_fields(db, project_id, register, values)
        obj = Model(project_id=project_id, number=await _next_number(db, project_id, cfg["prefix"]), **values)
        db.add(obj)
        await db.flush()
        await _set_links(db, project_id, cfg["item_type"], obj.id, data.meeting_ids)
        await db.commit()
        return (await _serialize(db, cfg, [obj]))[0]

    async def update_item(project_id: UUID, item_id: UUID, data: SchemaIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
        await _require(db, project_id, user, register, "edit")
        obj = (await db.execute(select(Model).where(Model.id == item_id, Model.project_id == project_id))).scalar_one_or_none()
        if not obj:
            raise HTTPException(404, "Item not found")
        values = data.model_dump(exclude={"meeting_ids"})
        await _validate_register_fields(db, project_id, register, values)
        for k, v in values.items():
            setattr(obj, k, v)
        await _set_links(db, project_id, cfg["item_type"], obj.id, data.meeting_ids)
        await db.commit()
        return (await _serialize(db, cfg, [obj]))[0]

    async def delete_item(project_id: UUID, item_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
        await _require(db, project_id, user, register, "edit")
        await db.execute(delete(PMMeetingLink).where(PMMeetingLink.item_type == cfg["item_type"], PMMeetingLink.item_id == item_id))
        await db.execute(delete(Model).where(Model.id == item_id, Model.project_id == project_id))
        await db.commit()

    router.add_api_route(base, list_items, methods=["GET"], response_model=List[SchemaOut], name=f"list_{register}")
    router.add_api_route(base, create_item, methods=["POST"], response_model=SchemaOut, status_code=201, name=f"create_{register}")
    router.add_api_route(base + "/{item_id}", update_item, methods=["PUT"], response_model=SchemaOut, name=f"update_{register}")
    router.add_api_route(base + "/{item_id}", delete_item, methods=["DELETE"], status_code=204, name=f"delete_{register}")


for _name, _cfg in REGISTERS.items():
    _register_routes(_name, _cfg)


# ─── Meetings ────────────────────────────────────────────────────────────────
async def _get_meeting(db: AsyncSession, project_id: UUID, meeting_id: UUID) -> PMMeeting:
    m = (await db.execute(select(PMMeeting).where(PMMeeting.id == meeting_id, PMMeeting.project_id == project_id))).scalar_one_or_none()
    if not m:
        raise HTTPException(404, "Meeting not found")
    return m


def _meeting_out(m: PMMeeting) -> MeetingResponse:
    return MeetingResponse.model_validate(m).model_copy(update={"has_transcript": bool(m.transcript_text)})


def _require_unlocked(m: PMMeeting):
    if m.status == "validated":
        raise HTTPException(409, "These meeting minutes are validated and can no longer be changed")


@router.get("/projects/{project_id}/meetings", response_model=List[MeetingSummary])
async def list_meetings(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "meetings")
    r = await db.execute(select(PMMeeting).where(PMMeeting.project_id == project_id).order_by(PMMeeting.meeting_date.desc().nullslast(), PMMeeting.number.desc()))
    return r.scalars().all()


# Meeting validations and deliverable approvals waiting on the caller, across every
# project — validators/approvers may have no section access of their own beyond being
# asked, so this is gated only on the module.
@router.get("/my-validations")
async def my_validations(db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    meetings = await db.execute(
        select(PMMeeting.id, PMMeeting.project_id, PMMeeting.number, PMMeeting.title, PMMeeting.meeting_date)
        .join(PMMeetingValidation, PMMeetingValidation.meeting_id == PMMeeting.id)
        .where(PMMeetingValidation.validator_email == user.email, PMMeetingValidation.status == "pending", PMMeeting.status == "in_review"))
    versions = await db.execute(
        select(PMDeliverable.id, PMDeliverable.project_id, PMDeliverable.number, PMDeliverable.name,
               PMDeliverableVersion.id.label("version_id"), PMDeliverableVersion.version_label)
        .join(PMDeliverableVersion, PMDeliverableVersion.deliverable_id == PMDeliverable.id)
        .join(PMDeliverableApproval, PMDeliverableApproval.version_id == PMDeliverableVersion.id)
        .where(PMDeliverableApproval.approver_email == user.email, PMDeliverableApproval.decision == "pending", PMDeliverableVersion.status == "in_approval"))
    return {
        "meetings": [dict(meeting_id=r.id, project_id=r.project_id, number=r.number, title=r.title, meeting_date=r.meeting_date) for r in meetings.all()],
        "deliverables": [dict(deliverable_id=r.id, project_id=r.project_id, number=r.number, name=r.name, version_id=r.version_id, version_label=r.version_label) for r in versions.all()],
    }


async def _require_meeting_read(db, project_id, meeting_id, user) -> PMMeeting:
    """Meeting readers: 'meetings' view access, or being one of its requested validators."""
    await _get_project(db, project_id)
    m = await _get_meeting(db, project_id, meeting_id)
    level = (await _sections_for(db, project_id, user))["meetings"]
    if level == "none" and not any(v.validator_email == user.email for v in m.validations):
        raise HTTPException(403, "No access to this meeting")
    return m


@router.get("/projects/{project_id}/meetings/{meeting_id}", response_model=MeetingResponse)
async def get_meeting(project_id: UUID, meeting_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    return _meeting_out(await _require_meeting_read(db, project_id, meeting_id, user))


@router.post("/projects/{project_id}/meetings", response_model=MeetingResponse, status_code=status.HTTP_201_CREATED)
async def create_meeting(project_id: UUID, data: MeetingIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "meetings", "edit")
    m = PMMeeting(project_id=project_id, number=await _next_number(db, project_id, "M"), created_by=user.email,
                  proposal=Proposal().model_dump(mode="json"), status="draft", **data.model_dump())
    db.add(m)
    await db.commit()
    await db.refresh(m, ["validations"])
    return _meeting_out(m)


@router.put("/projects/{project_id}/meetings/{meeting_id}", response_model=MeetingResponse)
async def update_meeting(project_id: UUID, meeting_id: UUID, data: MeetingIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "meetings", "edit")
    m = await _get_meeting(db, project_id, meeting_id)
    _require_unlocked(m)
    for k, v in data.model_dump().items():
        setattr(m, k, v)
    await db.commit()
    return _meeting_out(m)


@router.delete("/projects/{project_id}/meetings/{meeting_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_meeting(project_id: UUID, meeting_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    # Register entries created from this meeting stay; only their link to it goes (FK cascade).
    await _require(db, project_id, user, "meetings", "edit")
    await db.execute(delete(PMMeeting).where(PMMeeting.id == meeting_id, PMMeeting.project_id == project_id))
    await db.commit()


def _transcript_text(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if name.endswith((".txt", ".vtt", ".srt", ".md")):
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                return content.decode(enc)
            except UnicodeDecodeError:
                continue
    raise HTTPException(400, "Transcript must be a .docx, .txt, .vtt, .srt or .md file")


@router.post("/projects/{project_id}/meetings/{meeting_id}/transcript", response_model=MeetingResponse)
async def upload_transcript(project_id: UUID, meeting_id: UUID, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "meetings", "edit")
    m = await _get_meeting(db, project_id, meeting_id)
    _require_unlocked(m)
    content = await _read_upload(file)
    text_value = _transcript_text(file.filename, content).strip()
    if not text_value:
        raise HTTPException(400, "The transcript is empty")
    if len(text_value) > MAX_TRANSCRIPT_CHARS:
        raise HTTPException(413, "Transcript too long to process")
    m.transcript_ref = await upload_to_s3(f"projects/{project_id}/pm/meetings/{meeting_id}/{_safe_name(file.filename)}", content,
                                          file.content_type or "application/octet-stream")
    m.transcript_filename = file.filename
    m.transcript_text = text_value
    await db.commit()
    return _meeting_out(m)


_S = {"type": "string"}
MINUTES_SCHEMA = {
    "type": "object", "additionalProperties": False,
    "required": ["minutes", "actions", "decisions", "risks"],
    "properties": {
        "minutes": _S,
        "actions": {"type": "array", "items": {"type": "object", "additionalProperties": False,
            "required": ["existing_number", "title", "description", "owner", "due_date", "status", "comment"],
            "properties": {k: _S for k in ["existing_number", "title", "description", "owner", "due_date", "status", "comment"]}}},
        "decisions": {"type": "array", "items": {"type": "object", "additionalProperties": False,
            "required": ["existing_number", "title", "description", "decision_makers"],
            "properties": {"existing_number": _S, "title": _S, "description": _S, "decision_makers": {"type": "array", "items": _S}}}},
        "risks": {"type": "array", "items": {"type": "object", "additionalProperties": False,
            "required": ["existing_number", "title", "description", "mitigation", "impact", "probability", "status"],
            "properties": {k: _S for k in ["existing_number", "title", "description", "mitigation", "impact", "probability", "status"]}}},
    },
}


async def _call_claude(system: str, prompt: str, schema: dict = MINUTES_SCHEMA) -> dict:
    if not ANTHROPIC_API_KEY:
        raise HTTPException(500, "ANTHROPIC_API_KEY not configured")
    async with httpx.AsyncClient(timeout=600) as client:
        r = await client.post(
            "https://api.anthropic.com/v1/messages",
            headers={
                "x-api-key": ANTHROPIC_API_KEY,
                "anthropic-version": "2023-06-01",
                "anthropic-beta": "server-side-fallback-2026-07-01",
                "content-type": "application/json",
            },
            json={
                "model": MINUTES_MODEL, "max_tokens": 16000,
                "fallbacks": "default",
                "thinking": {"type": "adaptive"},
                "output_config": {"effort": "medium", "format": {"type": "json_schema", "schema": schema}},
                "system": system,
                "messages": [{"role": "user", "content": prompt}],
            },
        )
    if r.status_code != 200:
        raise HTTPException(502, f"Claude API error {r.status_code}: {r.text[:300]}")
    d = r.json()
    if d.get("stop_reason") == "refusal":
        raise HTTPException(502, "The AI declined to process this transcript")
    if d.get("stop_reason") == "max_tokens":
        raise HTTPException(502, "The AI output was cut off — the transcript may be too long")
    text_out = "".join(b.get("text", "") for b in d.get("content", []) if b.get("type") == "text")
    try:
        return json.loads(text_out)
    except json.JSONDecodeError:
        raise HTTPException(502, "Could not read the AI result")


MINUTES_SYSTEM = """You write the official minutes of project meetings from their transcripts, for a consulting company's project management office.
Write factual, neutral, professional minutes in the language used in the transcript. Never invent content that isn't supported by the transcript.

Return:
- minutes: the meeting minutes as plain text. Use lines starting with "# " for section headings (e.g. Context, Topics discussed, Next steps) and lines starting with "- " for bullet points. Do not repeat the action items, decisions and risks in full — they are listed separately.
- actions: every action item agreed in the meeting.
- decisions: every decision taken in the meeting.
- risks: every project risk raised in the meeting.

Linking to existing register entries: the user message lists the project's existing actions, decisions and risks with their numbers. When the meeting updates, closes, reschedules or re-discusses one of them, set existing_number to that entry's number and fill in its updated values; otherwise set existing_number to "".
Use "" for any unknown field. Dates are YYYY-MM-DD. Action status and risk impact/probability/status must be one of the allowed values listed in the user message."""

ACTIONS_SCHEMA = {
    "type": "object", "additionalProperties": False,
    "required": ["actions", "decisions", "risks"],
    "properties": {k: MINUTES_SCHEMA["properties"][k] for k in ("actions", "decisions", "risks")},
}

ACTIONS_SYSTEM = """You extract the action items, decisions and risks from a project meeting's finished, official minutes, for a consulting company's project management office.
The minutes are already written and final — do not rewrite, summarize or second-guess them, only extract structured entries from what they say. Never invent anything not supported by the text.

Return:
- actions: every action item agreed in the meeting.
- decisions: every decision taken in the meeting.
- risks: every project risk raised in the meeting.

Linking to existing register entries: the user message lists the project's existing actions, decisions and risks with their numbers. When the minutes update, close, reschedule or re-discuss one of them, set existing_number to that entry's number and fill in its updated values; otherwise set existing_number to "".
Use "" for any unknown field. Dates are YYYY-MM-DD. Action status and risk impact/probability/status must be one of the allowed values listed in the user message."""


async def _register_context(db: AsyncSession, project_id: UUID, s: PMSettings):
    """Existing open actions/decisions/risks plus the project's configured status/impact/
    probability values, formatted for a minutes-extraction prompt, and an existing_number ->
    id lookup so the AI's references can be resolved back to real register entries."""
    actions = (await db.execute(select(PMAction).where(PMAction.project_id == project_id).order_by(PMAction.number))).scalars().all()
    decisions = (await db.execute(select(PMDecision).where(PMDecision.project_id == project_id).order_by(PMDecision.number))).scalars().all()
    risks = (await db.execute(select(PMRisk).where(PMRisk.project_id == project_id).order_by(PMRisk.number))).scalars().all()

    existing = {
        "actions": [dict(number=a.number, title=a.title, description=a.description or "", owner=a.owner or "", status=a.status,
                         due_date=a.due_date.date().isoformat() if a.due_date else "") for a in actions if a.status != "Solved"],
        "decisions": [dict(number=d.number, title=d.title, description=d.description or "") for d in decisions],
        "risks": [dict(number=r.number, title=r.title, description=r.description or "", impact=r.impact, probability=r.probability,
                       status=r.status) for r in risks if r.status != "Closed"],
    }
    allowed = {
        "action_status": _action_statuses(s),
        "risk_impact": [l["level"] for l in s.impact_levels],
        "risk_probability": [l["level"] for l in s.probability_levels],
        "risk_status": ["Open", "Closed"],
    }
    by_number = {"actions": {a.number: a.id for a in actions}, "decisions": {d.number: d.id for d in decisions}, "risks": {r.number: r.id for r in risks}}
    return existing, allowed, by_number


def _proposal_items(result: dict, kind: str, by_number: dict) -> list:
    out = []
    for it in result.get(kind, []):
        it = {k: (v.strip() if isinstance(v, str) else v) for k, v in it.items()}
        it["existing_id"] = by_number[kind].get(it.pop("existing_number", ""))
        out.append({k: v for k, v in it.items() if v not in ("", None)} | {"title": it.get("title") or "(untitled)"})
    return out


@router.post("/projects/{project_id}/meetings/{meeting_id}/generate", response_model=MeetingResponse)
async def generate_minutes(project_id: UUID, meeting_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    proj = await _require(db, project_id, user, "meetings", "edit")
    m = await _get_meeting(db, project_id, meeting_id)
    _require_unlocked(m)
    if not m.transcript_text:
        raise HTTPException(400, "Upload the meeting transcript first")
    s = await _settings(db, project_id)
    existing, allowed, by_number = await _register_context(db, project_id, s)
    meeting_info = dict(project=proj.project_name, meeting_number=m.number, meeting_type=m.meeting_type or "", title=m.title,
                        date=m.meeting_date.date().isoformat() if m.meeting_date else "", attendees=m.attendees or [])
    prompt = (f"<meeting>\n{json.dumps(meeting_info, ensure_ascii=False)}\n</meeting>\n"
              f"<allowed_values>\n{json.dumps(allowed, ensure_ascii=False)}\n</allowed_values>\n"
              f"<existing_register_entries>\n{json.dumps(existing, ensure_ascii=False)}\n</existing_register_entries>\n"
              f"<transcript>\n{m.transcript_text}\n</transcript>")
    result = await _call_claude(MINUTES_SYSTEM, prompt)

    proposal = Proposal.model_validate({k: _proposal_items(result, k, by_number) for k in ("actions", "decisions", "risks")})
    m.minutes = result.get("minutes", "")
    m.proposal = proposal.model_dump(mode="json")
    m.status = "generated"
    await db.execute(delete(PMMeetingValidation).where(PMMeetingValidation.meeting_id == m.id))
    await db.commit()
    await db.refresh(m, ["validations"])
    return _meeting_out(m)


@router.post("/projects/{project_id}/meetings/{meeting_id}/generate-actions", response_model=MeetingResponse)
async def generate_actions_from_minutes(project_id: UUID, meeting_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    """For minutes that were written or pasted directly rather than generated from a
    transcript: extracts the action list, decision register and risk register from the
    already-final minutes text, without touching that text."""
    await _require(db, project_id, user, "meetings", "edit")
    m = await _get_meeting(db, project_id, meeting_id)
    _require_unlocked(m)
    if not (m.minutes or "").strip():
        raise HTTPException(400, "Write or generate the meeting minutes first")
    s = await _settings(db, project_id)
    existing, allowed, by_number = await _register_context(db, project_id, s)
    prompt = (f"<allowed_values>\n{json.dumps(allowed, ensure_ascii=False)}\n</allowed_values>\n"
              f"<existing_register_entries>\n{json.dumps(existing, ensure_ascii=False)}\n</existing_register_entries>\n"
              f"<minutes>\n{m.minutes}\n</minutes>")
    result = await _call_claude(ACTIONS_SYSTEM, prompt, ACTIONS_SCHEMA)

    proposal = Proposal.model_validate({k: _proposal_items(result, k, by_number) for k in ("actions", "decisions", "risks")})
    m.proposal = proposal.model_dump(mode="json")
    m.status = "generated"
    await db.execute(delete(PMMeetingValidation).where(PMMeetingValidation.meeting_id == m.id))
    await db.commit()
    await db.refresh(m, ["validations"])
    return _meeting_out(m)


@router.put("/projects/{project_id}/meetings/{meeting_id}/review", response_model=MeetingResponse)
async def save_review(project_id: UUID, meeting_id: UUID, data: MeetingReview, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "meetings", "edit")
    m = await _get_meeting(db, project_id, meeting_id)
    _require_unlocked(m)
    m.minutes = data.minutes
    m.proposal = data.proposal.model_dump(mode="json")
    # Any edit invalidates approvals already given on the previous content.
    if m.status == "in_review":
        await db.execute(delete(PMMeetingValidation).where(PMMeetingValidation.meeting_id == m.id))
    m.status = "generated"
    await db.commit()
    await db.refresh(m, ["validations"])
    return _meeting_out(m)


def _parse_date(value: str | None):
    if not value:
        return None
    try:
        return datetime.fromisoformat(value[:10])
    except ValueError:
        return None


async def _finalize_meeting(db: AsyncSession, m: PMMeeting):
    """Applies the reviewed proposal to the registers (update the linked entry, or create a
    new one) and links every touched entry to this meeting. Runs once — the status flip to
    'validated' locks the meeting against a second application."""
    s = await _settings(db, m.project_id)
    statuses = _action_statuses(s)
    impacts = [l["level"] for l in s.impact_levels]
    probabilities = [l["level"] for l in s.probability_levels]
    proposal = Proposal.model_validate(m.proposal)

    async def _target(model, existing_id, prefix):
        obj = None
        if existing_id:
            obj = (await db.execute(select(model).where(model.id == existing_id, model.project_id == m.project_id))).scalar_one_or_none()
        if obj is None:
            obj = model(project_id=m.project_id, number=await _next_number(db, m.project_id, prefix))
            db.add(obj)
        return obj

    for a in proposal.actions:
        obj = await _target(PMAction, a.existing_id, "A")
        is_new = obj.id is None
        obj.title = a.title
        for field in ("description", "owner", "comment"):
            if getattr(a, field):
                setattr(obj, field, getattr(a, field))
        if a.status in statuses:
            obj.status = a.status
        elif is_new:
            obj.status = "Open"
        if _parse_date(a.due_date):
            obj.due_date = _parse_date(a.due_date)
        if is_new:
            obj.opening_date = m.meeting_date or datetime.utcnow()
        if obj.status == "Solved" and not obj.closing_date:
            obj.closing_date = m.meeting_date or datetime.utcnow()
        await db.flush()
        a.applied_id = obj.id
        await _add_link(db, "action", obj.id, m.id)

    for d in proposal.decisions:
        obj = await _target(PMDecision, d.existing_id, "D")
        obj.title = d.title
        if d.description:
            obj.description = d.description
        if d.decision_makers:
            obj.decision_makers = d.decision_makers
        elif obj.decision_makers is None:
            obj.decision_makers = []
        obj.decision_date = m.meeting_date or obj.decision_date or datetime.utcnow()
        await db.flush()
        d.applied_id = obj.id
        await _add_link(db, "decision", obj.id, m.id)

    for r in proposal.risks:
        obj = await _target(PMRisk, r.existing_id, "R")
        is_new = obj.id is None
        obj.title = r.title
        for field in ("description", "mitigation"):
            if getattr(r, field):
                setattr(obj, field, getattr(r, field))
        if r.impact in impacts:
            obj.impact = r.impact
        elif is_new:
            obj.impact = impacts[len(impacts) // 2]
        if r.probability in probabilities:
            obj.probability = r.probability
        elif is_new:
            obj.probability = probabilities[len(probabilities) // 2]
        if r.status in ("Open", "Closed"):
            obj.status = r.status
        elif is_new:
            obj.status = "Open"
        await db.flush()
        r.applied_id = obj.id
        await _add_link(db, "risk", obj.id, m.id)

    m.proposal = proposal.model_dump(mode="json")
    m.status = "validated"
    m.validated_at = datetime.utcnow()


@router.post("/projects/{project_id}/meetings/{meeting_id}/request-validation", response_model=MeetingResponse)
async def request_validation(project_id: UUID, meeting_id: UUID, data: ValidationRequest, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    """Starts a fresh validation round. With no validators, the requester validates alone
    and the registers are updated immediately."""
    await _require(db, project_id, user, "meetings", "edit")
    m = await _get_meeting(db, project_id, meeting_id)
    _require_unlocked(m)
    if not (m.minutes or "").strip():
        raise HTTPException(400, "Generate or write the minutes before requesting validation")
    await db.execute(delete(PMMeetingValidation).where(PMMeetingValidation.meeting_id == m.id))
    validators = {}
    for v in data.validators:
        email = (v.email or "").strip().lower()
        if not email:
            raise HTTPException(400, f"Validator {v.name} has no email address")
        validators[email] = v.name
    if validators:
        for email, name in validators.items():
            db.add(PMMeetingValidation(meeting_id=m.id, validator_email=email, validator_name=name))
        m.status = "in_review"
    else:
        await _finalize_meeting(db, m)
    await db.commit()
    await db.refresh(m, ["validations"])
    return _meeting_out(m)


@router.post("/projects/{project_id}/meetings/{meeting_id}/validate", response_model=MeetingResponse)
async def decide_validation(project_id: UUID, meeting_id: UUID, data: ValidationDecision, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    m = await _require_meeting_read(db, project_id, meeting_id, user)
    if m.status != "in_review":
        raise HTTPException(409, "This meeting isn't awaiting validation")
    mine = next((v for v in m.validations if v.validator_email == user.email and v.status == "pending"), None)
    if not mine:
        raise HTTPException(403, "You have no pending validation on this meeting")
    mine.status = "approved" if data.approve else "rejected"
    mine.comment = data.comment
    mine.decided_at = datetime.utcnow()
    if not data.approve:
        m.status = "generated"  # back to the reviewer, who edits and requests a new round
    elif all(v.status == "approved" for v in m.validations):
        await _finalize_meeting(db, m)
    await db.commit()
    await db.refresh(m, ["validations"])
    return _meeting_out(m)


# ─── Meeting minutes export (.docx) ──────────────────────────────────────────
PLACEHOLDERS = ("{{PROJECT_NAME}}", "{{PROJECT_NUMBER}}", "{{MEETING_NUMBER}}", "{{MEETING_TITLE}}", "{{MEETING_TYPE}}", "{{MEETING_DATE}}")


def _replace_placeholders(doc, values: dict):
    def _in_paragraphs(paragraphs):
        for p in paragraphs:
            if "{{" not in p.text:
                continue
            new_text = p.text
            for k, v in values.items():
                new_text = new_text.replace(k, v)
            if new_text != p.text and p.runs:
                p.runs[0].text = new_text
                for run in p.runs[1:]:
                    run.text = ""

    def _in_container(c):
        _in_paragraphs(c.paragraphs)
        for t in c.tables:
            for row in t.rows:
                for cell in row.cells:
                    _in_container(cell)

    _in_container(doc)
    for section in doc.sections:
        _in_container(section.header)
        _in_container(section.footer)


def _add_table(doc, headers: list, rows: list):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value or ""


@router.get("/projects/{project_id}/meetings/{meeting_id}/export")
async def export_minutes(project_id: UUID, meeting_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    from docx import Document
    from docx.shared import Cm

    proj = await _require(db, project_id, user, "meetings")
    m = await _get_meeting(db, project_id, meeting_id)
    if m.status != "validated":
        raise HTTPException(409, "Only validated meeting minutes can be exported")
    s = await _settings(db, project_id)
    date_str = m.meeting_date.strftime("%d/%m/%Y") if m.meeting_date else ""
    values = dict(zip(PLACEHOLDERS, [proj.project_name, proj.project_number or "", m.number, m.title, m.meeting_type or "", date_str]))

    template = (await db.execute(select(PMTemplate).where(PMTemplate.project_id == project_id, PMTemplate.template_type == "meeting_minutes",
                                                          PMTemplate.filename.ilike("%.docx")).order_by(PMTemplate.created_at.desc()))).scalars().first()
    if template:
        doc = Document(io.BytesIO(await _s3_get(template.file_ref)))
        _replace_placeholders(doc, values)
    else:
        doc = Document()
        if s.customer_logo_ref:
            try:
                header_p = doc.sections[0].header.paragraphs[0]
                header_p.add_run().add_picture(io.BytesIO(await _s3_get(s.customer_logo_ref)), height=Cm(1.5))
            except Exception as e:  # a broken logo shouldn't block the export
                print(f"PM export: logo skipped: {e}")
        doc.add_heading(f"Meeting minutes — {m.title}", level=0)
        doc.add_paragraph(f"{proj.project_name} ({proj.project_number or ''}) · {m.number} · {m.meeting_type or ''} · {date_str}")

    doc.add_heading("Attendees", level=1)
    for a in m.attendees or []:
        doc.add_paragraph(f"{a.get('name', '')}{' — ' + a['email'] if a.get('email') else ''}", style="List Bullet")

    doc.add_heading("Minutes", level=1)
    for line in (m.minutes or "").splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line.strip())

    proposal = Proposal.model_validate(m.proposal)
    numbers = {}
    for model, items in ((PMAction, proposal.actions), (PMDecision, proposal.decisions), (PMRisk, proposal.risks)):
        ids = [i.applied_id for i in items if i.applied_id]
        if ids:
            numbers.update({row.id: row.number for row in (await db.execute(select(model.id, model.number).where(model.id.in_(ids)))).all()})
    if proposal.decisions:
        doc.add_heading("Decisions", level=1)
        _add_table(doc, ["#", "Decision", "Description", "Decision makers"],
                   [[numbers.get(d.applied_id, ""), d.title, d.description, ", ".join(d.decision_makers)] for d in proposal.decisions])
    if proposal.actions:
        doc.add_heading("Action items", level=1)
        _add_table(doc, ["#", "Action", "Owner", "Due date", "Status"],
                   [[numbers.get(a.applied_id, ""), a.title + (f"\n{a.description}" if a.description else ""), a.owner, a.due_date, a.status] for a in proposal.actions])
    if proposal.risks:
        doc.add_heading("Risks", level=1)
        _add_table(doc, ["#", "Risk", "Mitigation", "Impact", "Probability"],
                   [[numbers.get(r.applied_id, ""), r.title, r.mitigation, r.impact, r.probability] for r in proposal.risks])
    if m.validations:
        doc.add_heading("Validation", level=1)
        _add_table(doc, ["Validator", "Status", "Date", "Comment"],
                   [[v.validator_name or v.validator_email, v.status, v.decided_at.strftime("%d/%m/%Y") if v.decided_at else "", v.comment] for v in m.validations])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = _safe_name(f"{proj.project_number or 'project'}_{m.number}_minutes.docx")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# ─── Deliverables ────────────────────────────────────────────────────────────
async def _get_deliverable(db, project_id, deliverable_id) -> PMDeliverable:
    d = (await db.execute(select(PMDeliverable).where(PMDeliverable.id == deliverable_id, PMDeliverable.project_id == project_id))).scalar_one_or_none()
    if not d:
        raise HTTPException(404, "Deliverable not found")
    return d


async def _get_version(db, deliverable: PMDeliverable, version_id) -> PMDeliverableVersion:
    v = next((v for v in deliverable.versions if v.id == version_id), None)
    if not v:
        raise HTTPException(404, "Version not found")
    return v


async def _reload_deliverable(db, d: PMDeliverable) -> PMDeliverable:
    return (await db.execute(select(PMDeliverable).where(PMDeliverable.id == d.id).execution_options(populate_existing=True))).scalar_one()


def _deliverable_values(data: DeliverableIn) -> dict:
    v = data.model_dump(mode="json")
    for key in ("approvers", "contributors"):
        v[key] = [{**p, "email": (p.get("email") or "").strip().lower() or None} for p in v[key]]
    if v["owner"]:
        v["owner"]["email"] = (v["owner"].get("email") or "").strip().lower() or None
    return v


@router.get("/projects/{project_id}/deliverables", response_model=List[DeliverableResponse])
async def list_deliverables(project_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables")
    r = await db.execute(select(PMDeliverable).where(PMDeliverable.project_id == project_id).order_by(PMDeliverable.number))
    return r.scalars().all()


@router.post("/projects/{project_id}/deliverables", response_model=DeliverableResponse, status_code=status.HTTP_201_CREATED)
async def add_deliverable(project_id: UUID, data: DeliverableIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables", "edit")
    d = PMDeliverable(project_id=project_id, number=await _next_number(db, project_id, "DEL"), **_deliverable_values(data))
    db.add(d)
    await db.commit()
    return await _reload_deliverable(db, d)


@router.put("/projects/{project_id}/deliverables/{deliverable_id}", response_model=DeliverableResponse)
async def update_deliverable(project_id: UUID, deliverable_id: UUID, data: DeliverableIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables", "edit")
    d = await _get_deliverable(db, project_id, deliverable_id)
    for k, v in _deliverable_values(data).items():
        setattr(d, k, v)
    await db.commit()
    return await _reload_deliverable(db, d)


@router.delete("/projects/{project_id}/deliverables/{deliverable_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_deliverable(project_id: UUID, deliverable_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables", "edit")
    await db.execute(delete(PMDeliverable).where(PMDeliverable.id == deliverable_id, PMDeliverable.project_id == project_id))
    await db.commit()


@router.post("/projects/{project_id}/deliverables/{deliverable_id}/versions", response_model=DeliverableResponse, status_code=status.HTTP_201_CREATED)
async def add_version(project_id: UUID, deliverable_id: UUID, data: VersionIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables", "edit")
    d = await _get_deliverable(db, project_id, deliverable_id)
    if any(v.version_label == data.version_label.strip() for v in d.versions):
        raise HTTPException(409, "This version already exists")
    db.add(PMDeliverableVersion(deliverable_id=d.id, **{**data.model_dump(), "version_label": data.version_label.strip()}))
    await db.commit()
    return await _reload_deliverable(db, d)


@router.put("/projects/{project_id}/deliverables/{deliverable_id}/versions/{version_id}", response_model=DeliverableResponse)
async def update_version(project_id: UUID, deliverable_id: UUID, version_id: UUID, data: VersionIn, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables", "edit")
    d = await _get_deliverable(db, project_id, deliverable_id)
    v = await _get_version(db, d, version_id)
    if v.status in ("in_approval", "approved"):
        raise HTTPException(409, "A version in approval or approved can't be edited — create a new version")
    for k, val in data.model_dump().items():
        setattr(v, k, val)
    await db.commit()
    return await _reload_deliverable(db, d)


@router.delete("/projects/{project_id}/deliverables/{deliverable_id}/versions/{version_id}", response_model=DeliverableResponse)
async def delete_version(project_id: UUID, deliverable_id: UUID, version_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables", "edit")
    d = await _get_deliverable(db, project_id, deliverable_id)
    v = await _get_version(db, d, version_id)
    if v.status == "approved":
        raise HTTPException(409, "Approved versions are kept as the approval record")
    await db.delete(v)
    await db.commit()
    return await _reload_deliverable(db, d)


@router.post("/projects/{project_id}/deliverables/{deliverable_id}/versions/{version_id}/submit", response_model=DeliverableResponse)
async def submit_version(project_id: UUID, deliverable_id: UUID, version_id: UUID, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    await _require(db, project_id, user, "deliverables", "edit")
    d = await _get_deliverable(db, project_id, deliverable_id)
    v = await _get_version(db, d, version_id)
    if v.status not in ("draft", "rejected"):
        raise HTTPException(409, "Only draft or rejected versions can be submitted for approval")
    approvers = {p["email"]: p.get("name") for p in (d.approvers or []) if p.get("email")}
    if not approvers:
        raise HTTPException(400, "Add approvers with an email address to the deliverable first")
    await db.execute(delete(PMDeliverableApproval).where(PMDeliverableApproval.version_id == v.id))
    for email, name in approvers.items():
        db.add(PMDeliverableApproval(version_id=v.id, approver_email=email, approver_name=name))
    v.status = "in_approval"
    v.submitted_at = datetime.utcnow()
    await db.commit()
    return await _reload_deliverable(db, d)


@router.post("/projects/{project_id}/deliverables/{deliverable_id}/versions/{version_id}/approve", response_model=DeliverableResponse)
async def decide_version(project_id: UUID, deliverable_id: UUID, version_id: UUID, data: ValidationDecision, db: AsyncSession = Depends(get_db), user: PMUser = Depends(pm_user)):
    # Approvers need no 'deliverables' access beyond being named on this version.
    await _get_project(db, project_id)
    d = await _get_deliverable(db, project_id, deliverable_id)
    v = await _get_version(db, d, version_id)
    if v.status != "in_approval":
        raise HTTPException(409, "This version isn't awaiting approval")
    mine = next((a for a in v.approvals if a.approver_email == user.email and a.decision == "pending"), None)
    if not mine:
        raise HTTPException(403, "You have no pending approval on this version")
    mine.decision = "approved" if data.approve else "rejected"
    mine.comment = data.comment
    mine.decided_at = datetime.utcnow()
    if not data.approve:
        v.status = "rejected"
    elif all(a.decision == "approved" for a in v.approvals):
        v.status = "approved"
        v.approved_at = datetime.utcnow()
        if v.document_url:
            d.document_url = v.document_url
    await db.commit()
    return await _reload_deliverable(db, d)
